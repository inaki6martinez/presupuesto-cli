"""El ajuste conserva el saldo al categorizar, dividir o cancelar."""

from decimal import Decimal

import openpyxl
import pytest
from click.testing import CliRunner

from presupuesto import cmd_actualizar as modulo
from presupuesto.tui_categorizar import TUICategorizacion
from presupuesto.tui_dividir import TUIDividir
from presupuesto.tui_revision import TUIRevisionFinal
from tests.test_integracion import _crear_xlsx_test


@pytest.mark.parametrize("modo,nuevo,filas", [
    ("editar", "70", 1),
    ("dividir", "70", 2),
    ("dividir", "130", 2),
    ("cancelar", "70", 0),
    ("volver", "70", 1),
    ("suma_incorrecta", "70", 0),
    ("sin_diferencia", "100", 0),
])
def test_actualizar_categoriza_y_conserva_saldo(tmp_path, monkeypatch, modo, nuevo, filas):
    ruta = tmp_path / "presupuesto.xlsx"
    _crear_xlsx_test(ruta)
    wb = openpyxl.load_workbook(ruta)
    wb["Datos"].append([2026, "Sep", "Finanzas", "Balance", "", "", 100,
                        "", "", "Cuenta Ocio", "N26", "Activos liquidos", "Real"])
    wb.save(ruta)
    wb.close()
    monkeypatch.setattr("presupuesto.config.cargar_config", lambda: {"archivo_presupuesto": str(ruta)})
    revisiones = []
    monkeypatch.setattr(modulo, "_pedir_registrar_revision", lambda *args: revisiones.append(args[1]))
    balances_finales = {}
    selecciones = iter([("Cuenta Ocio", "N26", "Activos liquidos"), None])

    def seleccionar(cuentas, balances):
        balances_finales.update(balances)
        return next(selecciones)

    monkeypatch.setattr(modulo, "_tui_seleccionar_cuenta", seleccionar)
    categorias = iter([
        {"categoria1": "Gastos Personales", "categoria2": "Bebé", "tipo_gasto": "Optimizable"},
        {"categoria1": "Finanzas", "categoria2": "Balance"},
    ])
    monkeypatch.setattr(TUICategorizacion, "run", lambda self: next(categorias))
    signo = Decimal("-1") if Decimal(nuevo) < 100 else Decimal("1")
    monkeypatch.setattr(TUIDividir, "run", lambda self: [
        (signo * 10, "Compra bebé"), (signo * 20, "Resto"),
    ])
    visitas = []

    def revisar(self):
        visitas.append(True)
        assert self._movs[0].cuenta == "Cuenta Ocio"
        if modo == "cancelar":
            return False
        if modo == "volver" and len(visitas) == 1:
            return "volver"
        if modo == "dividir":
            self._dividir_movimiento(0)
        elif modo == "suma_incorrecta":
            self._movs[0].importe = Decimal("999")
        else:
            self._editar_movimiento(0)
        return True

    monkeypatch.setattr(TUIRevisionFinal, "run", revisar)
    resultado = CliRunner().invoke(modulo.cmd_actualizar, input=nuevo + "\n")
    assert resultado.exit_code == 0, resultado.output
    assert balances_finales["Cuenta Ocio"] == (Decimal(nuevo) if filas else Decimal("100"))
    assert len(revisiones) == (1 if filas or modo == "sin_diferencia" else 0)
    wb = openpyxl.load_workbook(ruta, read_only=True)
    datos = list(wb["Datos"].iter_rows(min_row=3, values_only=True))
    wb.close()
    assert len(datos) == filas
    assert len(list(tmp_path.glob("*backup*.xlsx"))) == (1 if filas else 0)
    if filas:
        assert sum(Decimal(str(r[6])) for r in datos) == Decimal(nuevo) - 100
        assert datos[0][2:4] == ("Gastos Personales", "Bebé")
        assert all(r[9:13] == ("Cuenta Ocio", "N26", "Activos liquidos", "Real") for r in datos)
        if filas == 2:
            assert datos[1][2:4] == ("Finanzas", "Balance")
