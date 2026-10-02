"""Regresiones del flujo de revisión, con libros pequeños y temporales."""
from dataclasses import replace
from decimal import Decimal

import openpyxl
from rich.console import Console

from presupuesto.categorizar import MovimientoCategorizado
from presupuesto.escritor import leer_numero, eliminar_filas


def movimiento(**cambios):
    base = MovimientoCategorizado(2026, "Sep", "Ocio", "Compra", "", "",
                                 Decimal("-50"), "", "Fijos", "Cuenta", "Banco", "Activo")
    return replace(base, **cambios)


def libro(ruta):
    wb = openpyxl.Workbook()
    wb.active.title = "Datos"
    wb.active.append(["Cabecera"] * 13)
    wb.create_sheet("Maestro").append(["Cabecera"] * 11)
    wb.create_sheet("Claves").append(["Cuenta", "Banco", "Tipo"])
    wb["Claves"].append(["Cuenta", "Banco", "Activo"])
    wb["Datos"].append([2026, "Sep", "Ocio", "Compra", "", "", "=-50", "", "Fijos",
                        "Cuenta", '=VLOOKUP(J2,Claves!$A:$C,2,0)',
                        '=VLOOKUP(J2,Claves!$A:$C,3,0)', "Real"])
    wb.save(ruta)
    wb.close()


def test_formulas_lectura_cierre_y_borrado(tmp_path):
    from io import StringIO
    from presupuesto.duplicados import detectar_duplicados
    from presupuesto.cmd_actualizar import leer_balances
    from presupuesto.cmd_cerrar import _analizar, _plan, _ejecutar
    from presupuesto.cmd_vista import _cmd_vista_mes, _leer_datos

    ruta = tmp_path / "datos.xlsx"
    libro(ruta)
    assert detectar_duplicados([movimiento()], ruta)[0][1] == 2
    assert leer_balances(ruta)["Cuenta"] == Decimal("-50")
    salida = StringIO()
    _cmd_vista_mes(Console(file=salida, width=120), ruta, "Sep", 2026)
    assert "-50" in salida.getvalue()
    datos = _analizar(ruta)
    _ejecutar(ruta, 2026, _plan(datos, 2026, 10))
    wb = openpyxl.load_workbook(ruta)
    assert leer_numero(wb["Datos"].cell(3, 7).value) == -50
    eliminar_filas(wb["Datos"], [2])
    assert "J2," in wb["Datos"].cell(2, 11).value
    wb["Datos"].cell(2, 3, "Finanzas")
    wb["Datos"].cell(2, 4, "Balance")
    wb.save(ruta)
    wb.close()
    entradas, *_ = _leer_datos(ruta, [(2027, "Sep")], incluir_balance=True)
    assert any(f.cat2 == "Balance" for f in entradas)


def test_exportacion_y_origen(tmp_path):
    import csv
    from presupuesto.cli import _exportar_csv
    from presupuesto.agrupador import agrupar_movimientos

    origen = [{"fecha": "2026-09-10", "cuenta": "Cuenta", "concepto": "Compra"}]
    original = movimiento(importe=Decimal("-30"), originales=origen, fuente="manual")
    agrupado = agrupar_movimientos([original])[0]
    assert agrupado.originales == origen and agrupado.fuente == "manual"
    partes = [replace(agrupado, importe=Decimal("-10"), categoria1="Salud"),
              replace(agrupado, importe=Decimal("-20"))]
    ruta = tmp_path / "exportado.csv"
    _exportar_csv(partes, str(ruta))
    with ruta.open(encoding="utf-8-sig") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    assert sum(Decimal(r["importe"]) for r in filas) == -30
    assert filas[0]["categoria1"] == "Salud"


def test_volver_a_duplicados_conserva_edicion(tmp_path, monkeypatch):
    from datetime import date
    from importlib import import_module
    from unittest.mock import Mock
    from click.testing import CliRunner
    from presupuesto.parsers.base import MovimientoCrudo
    from presupuesto.duplicados import GestorMarcadores, GestorRevisiones
    from presupuesto.tui_revision import TUIRevisionFinal, TUIRevisionDuplicados

    cli = import_module("presupuesto.cli")
    ruta = tmp_path / "datos.xlsx"
    libro(ruta)
    crudos = [MovimientoCrudo(date(2026, 9, 30), "Duplicado", Decimal("-50"), "Duplicado"),
              MovimientoCrudo(date(2026, 9, 10), "Nuevo", Decimal("-5"), "Nuevo")]
    parser = Mock()
    parser.parsear.return_value = crudos
    categorizador = Mock()
    categorizador.categorizar.side_effect = lambda m, c: movimiento(
        importe=m.importe, confianza="alta", concepto_original=m.concepto,
        originales=[{"fecha": m.fecha.isoformat(), "cuenta": c, "concepto": m.concepto}])
    visitas, escritos = [], []

    def revisar(self):
        visitas.append(self._movs[0].categoria1)
        if len(visitas) == 1:
            self._movs[0] = replace(self._movs[0], categoria1="Salud")
            return "volver"
        return True

    escritor = Mock()
    escritor.escribir.side_effect = lambda ms: escritos.extend(ms) or len(ms)
    monkeypatch.setattr(cli, "_obtener_ruta_xlsx", lambda: ruta)
    monkeypatch.setattr(cli, "_crear_gestor_reglas", Mock)
    monkeypatch.setattr(cli, "_obtener_parser_y_banco", lambda *args: (parser, "n26"))
    monkeypatch.setattr("presupuesto.config.cargar_config", lambda: {"cuentas_defecto": {"n26": "Cuenta"}})
    monkeypatch.setattr("presupuesto.categorizar.Categorizador", lambda *args: categorizador)
    marcadores = GestorMarcadores(tmp_path / "marcadores.json")
    monkeypatch.setattr("presupuesto.duplicados.GestorMarcadores", lambda: marcadores)
    monkeypatch.setattr("presupuesto.duplicados.GestorRevisiones", lambda: GestorRevisiones(tmp_path / "revisiones.json"))
    monkeypatch.setattr("presupuesto.escritor.EscritorDatos", lambda *args: escritor)
    monkeypatch.setattr(TUIRevisionDuplicados, "run", lambda self: {0})
    monkeypatch.setattr(TUIRevisionFinal, "run", revisar)
    resultado = CliRunner().invoke(cli.cmd_importar, [str(ruta), "--no-interactivo"])
    assert resultado.exit_code == 0, resultado.exception
    assert visitas == ["Ocio", "Salud"]
    assert escritos[0].categoria1 == "Salud"
